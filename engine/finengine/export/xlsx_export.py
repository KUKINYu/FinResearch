"""Excel 导出：同行对比表（openpyxl）。"""

import datetime
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HEADER_FILL = PatternFill("solid", fgColor="1B3A6B")
HEADER_FONT = Font(name="微软雅黑", color="FFFFFF", bold=True, size=11)
CELL_FONT = Font(name="微软雅黑", size=10)


def export_comparison_xlsx(companies: list[dict], note: str = "") -> bytes:
    """对比表：行 = 指标，列 = 公司。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "同行对比"

    indicators = [
        ("营业收入（万元）", lambda c: c["financials"].get("营业收入")),
        ("归母净利润（万元）", lambda c: c["financials"].get("归母净利润")),
        ("毛利率（%）", lambda c: c["financials"].get("毛利率")),
        ("净利率（%）", lambda c: c["financials"].get("净利率")),
        ("ROE（%）", lambda c: c["financials"].get("ROE")),
        ("研发投入（万元）", lambda c: c["financials"].get("研发投入")),
        ("PE（动态）", lambda c: c.get("valuation", {}).get("pe")),
        ("PB", lambda c: c.get("valuation", {}).get("pb")),
        ("总市值（亿元）", lambda c: (c.get("valuation", {}).get("market_cap") or 0) / 1e8),
    ]

    year = companies[0].get("year", "—") if companies else "—"
    ws.append([f"数据年度：{year}"])
    ws.append(["指标"] + [c["name"] for c in companies])
    for label, getter in indicators:
        row = [label]
        for c in companies:
            v = getter(c)
            row.append(round(v, 2) if isinstance(v, float) else (v if v is not None else "—"))
        ws.append(row)

    # 样式
    for cell in ws[2]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=3):
        for cell in row:
            cell.font = CELL_FONT
    ws.column_dimensions["A"].width = 22
    for col in ws.columns:
        col_letter = col[0].column_letter
        if col_letter != "A":
            ws.column_dimensions[col_letter].width = 16

    if note:
        ws.append([])
        ws.append([f"说明：{note} 生成时间：{datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}"])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
