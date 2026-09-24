"""Word 导出：异常清单与财务指标表（python-docx）。"""

import datetime
import io
import json

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

NAVY = RGBColor(0x1B, 0x3A, 0x6B)
SEVERITY_TEXT = {"high": "高", "medium": "中", "low": "低"}


def _set_cn_font(run, size: float = 10.5, bold: bool = False, color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def _title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_cn_font(p.add_run(text), size=16, bold=True, color=NAVY)


def _meta(doc: Document, lines: list[str]) -> None:
    for line in lines:
        p = doc.add_paragraph()
        _set_cn_font(p.add_run(line), size=9, color=RGBColor(0x7A, 0x8A, 0xA0))


def export_anomalies_docx(project: dict, anomalies: list[dict]) -> bytes:
    """异常清单：结论 + 计算过程 + 数据明细表（含出处页码）。"""
    doc = Document()
    _title(doc, f"{project.get('name', '项目')} — 财务异常清单")
    _meta(doc, [
        f"标的公司：{project.get('company_name') or '—'}",
        f"生成时间：{datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}",
        f"异常数量：{len(anomalies)} 条",
    ])
    doc.add_paragraph()

    if not anomalies:
        p = doc.add_paragraph()
        _set_cn_font(p.add_run("未发现触发阈值的异常。"), size=10.5)
        return _to_bytes(doc)

    for i, a in enumerate(anomalies, 1):
        # 标题
        p = doc.add_paragraph()
        _set_cn_font(p.add_run(f"{i}. [{SEVERITY_TEXT.get(a['severity'], a['severity'])}] {a['title']}"),
                     size=12, bold=True, color=NAVY)
        # 结论与计算过程
        p = doc.add_paragraph()
        _set_cn_font(p.add_run(a["description"]), size=10.5)
        # 数据明细表
        points = a.get("data_points") or []
        if points:
            table = doc.add_table(rows=1, cols=4)
            table.style = "Table Grid"
            headers = ["指标", "期间", "数值", "出处"]
            for j, h in enumerate(headers):
                cell = table.rows[0].cells[j]
                _set_cn_font(cell.paragraphs[0].add_run(h), bold=True)
            for pt_ in points:
                row = table.add_row().cells
                _set_cn_font(row[0].paragraphs[0].add_run(str(pt_.get("indicator", ""))))
                _set_cn_font(row[1].paragraphs[0].add_run(str(pt_.get("period", ""))))
                _set_cn_font(row[2].paragraphs[0].add_run(
                    f"{pt_.get('value', ''):,}{pt_.get('unit', '')}"))
                _set_cn_font(row[3].paragraphs[0].add_run(
                    f"第{pt_['source_page']}页" if pt_.get("source_page") else "—"))
        doc.add_paragraph()
    return _to_bytes(doc)


def _period_sort_key(period: str) -> tuple:
    import re

    m = re.search(r"(\d{4})", period)
    year = int(m.group(1)) if m else 0
    interim = 1 if "1-6月" in period else 0
    return (year, -interim)


def export_indicators_docx(project: dict, indicators: list[dict]) -> bytes:
    """财务指标表：指标 × 期间矩阵（含出处页码标注）。"""
    doc = Document()
    _title(doc, f"{project.get('name', '项目')} — 财务指标表")
    _meta(doc, [
        f"标的公司：{project.get('company_name') or '—'}",
        f"生成时间：{datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}",
    ])
    doc.add_paragraph()

    # 组织：指标 → 期间 → (值, 出处)
    by_name: dict[str, dict[str, dict]] = {}
    for ind in indicators:
        by_name.setdefault(ind["name"], {})[ind["period"]] = ind
    periods = sorted({ind["period"] for ind in indicators}, key=_period_sort_key, reverse=True)
    names = sorted(by_name.keys())

    table = doc.add_table(rows=1, cols=len(periods) + 1)
    table.style = "Table Grid"
    _set_cn_font(table.rows[0].cells[0].paragraphs[0].add_run("指标"), bold=True)
    for j, period in enumerate(periods, 1):
        _set_cn_font(table.rows[0].cells[j].paragraphs[0].add_run(period), bold=True)
    for name in names:
        cells = table.add_row().cells
        _set_cn_font(cells[0].paragraphs[0].add_run(name), bold=True)
        for j, period in enumerate(periods, 1):
            ind = by_name[name].get(period)
            if ind is None:
                text = "—"
            else:
                text = f"{ind['value']:,.2f}{ind['unit']}"
                if ind.get("source_page"):
                    text += f"（第{ind['source_page']}页）"
            _set_cn_font(cells[j].paragraphs[0].add_run(text), size=9)
    return _to_bytes(doc)


def _to_bytes(doc: Document) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
